import re
import shutil
import subprocess
import time
from pathlib import Path
import openpyxl

BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR / "input" / "job-tracker.xlsx"
RESUMES_DIR = BASE_DIR / "resumes"
PDFS_DIR = BASE_DIR / "pdfs"
LOGS_DIR = BASE_DIR / "logs"
LOG_FILE = LOGS_DIR / "resume-generation.log"

SKIP_UNCHANGED = True
REQUIRE_ONE_PAGE = True
LATEX_TIMEOUT = 60

for d in [INPUT_FILE.parent, RESUMES_DIR, PDFS_DIR, LOGS_DIR]:
    d.mkdir(parents=True, exist_ok=True)

def log(message):
    line = f"[{time.strftime('%Y-%m-%d %H:%M:%S')}] {message}"
    print(line)
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(line + "\n")

def find_column(ws, name):
    for cell in ws[1]:
        if cell.value == name:
            return cell.column
    return None

def tex_name(rank):
    return f"Priya_Resume_3yrs({rank}).tex"

def pdf_name(rank):
    return f"Priya_Resume_3yrs({rank}).pdf"

def clean_tex(tex):
    if not tex:
        return ""
    tex = str(tex).strip()
    tex = re.sub(r"^```(?:latex|tex)?\s*", "", tex, flags=re.I)
    tex = re.sub(r"\s*```\s*$", "", tex, flags=re.I)
    return tex.strip()

def validate_tex(tex, rank):
    if not tex:
        raise ValueError(f"Empty Tailored Resume .tex for rank {rank}")
    if not tex.startswith(r"\documentclass"):
        raise ValueError(f"Resume for rank {rank} does not start with \\documentclass")
    if r"\begin{document}" not in tex:
        raise ValueError(f"Resume for rank {rank} is missing \\begin{{document}}")
    if r"\end{document}" not in tex:
        raise ValueError(f"Resume for rank {rank} is missing \\end{{document}}")

def write_tex(rank, content):
    content = clean_tex(content)
    validate_tex(content, rank)
    path = RESUMES_DIR / tex_name(rank)
    path.write_text(content, encoding="utf-8")
    return path

def page_count(pdf):
    pdfinfo = shutil.which("pdfinfo")
    if not pdfinfo:
        return None
    result = subprocess.run([pdfinfo, str(pdf)], capture_output=True, text=True, timeout=20)
    if result.returncode != 0:
        raise RuntimeError(f"pdfinfo failed for {pdf.name}")
    match = re.search(r"^Pages:\s+(\d+)", result.stdout, re.M)
    if not match:
        raise RuntimeError(f"Could not determine page count for {pdf.name}")
    return int(match.group(1))

def compile_pdf(tex_path, rank):
    pdflatex = shutil.which("pdflatex")
    if not pdflatex:
        raise RuntimeError("pdflatex was not found. Install MacTeX/TeX Live and ensure pdflatex is in PATH.")

    pdf = PDFS_DIR / pdf_name(rank)

    if SKIP_UNCHANGED and pdf.exists() and pdf.stat().st_mtime >= tex_path.stat().st_mtime:
        log(f"Skipping unchanged resume: {pdf.name}")
        return pdf, True

    if pdf.exists():
        pdf.unlink()

    command = [
        pdflatex, "-interaction=nonstopmode", "-halt-on-error",
        "-output-directory", str(PDFS_DIR), str(tex_path)
    ]
    log(f"Compiling {tex_path.name}")
    result = subprocess.run(command, capture_output=True, text=True, timeout=LATEX_TIMEOUT)

    if result.returncode != 0:
        latex_log = LOGS_DIR / f"latex_{rank}.log"
        latex_log.write_text(result.stdout, encoding="utf-8")
        raise RuntimeError(f"LaTeX compilation failed for rank {rank}. See {latex_log}")

    if not pdf.exists():
        raise RuntimeError(f"PDF was not created: {pdf}")

    log(f"Created PDF: {pdf.name}")

    if REQUIRE_ONE_PAGE:
        pages = page_count(pdf)
        if pages is None:
            log("WARNING: pdfinfo not installed. Skipping page-count validation.")
        elif pages != 1:
            raise RuntimeError(f"{pdf.name} contains {pages} pages. Expected exactly 1 page.")
        else:
            log(f"PDF validated: {pdf.name} = 1 page")

    return pdf, False

def cleanup_aux():
    for p in PDFS_DIR.iterdir():
        if p.is_file() and p.suffix in {".aux", ".log", ".out"}:
            try:
                p.unlink()
            except OSError:
                pass

def main():
    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"Excel file not found: {INPUT_FILE}")

    if not shutil.which("pdflatex"):
        raise RuntimeError("pdflatex was not found. Run `which pdflatex` and make sure MacTeX is installed.")

    wb = openpyxl.load_workbook(INPUT_FILE, data_only=False)
    ws = wb.active

    rank_col = find_column(ws, "Rank")
    tex_col = find_column(ws, "Tailored Resume .tex")

    if not rank_col:
        raise RuntimeError("Required Excel column missing: Rank")
    if not tex_col:
        raise RuntimeError("Required Excel column missing: Tailored Resume .tex")

    total = generated = skipped = failed = 0

    log(f"Using pdflatex: {shutil.which('pdflatex')}")
    log(f"Loading Excel: {INPUT_FILE}")

    for row in range(2, ws.max_row + 1):
        rank = ws.cell(row, rank_col).value
        if rank is None:
            continue

        total += 1
        try:
            log("=" * 56)
            log(f"Processing rank {rank}")
            tex = write_tex(rank, ws.cell(row, tex_col).value)
            log(f"Created TEX: {tex.name}")
            _, was_skipped = compile_pdf(tex, rank)
            if was_skipped:
                skipped += 1
            else:
                generated += 1
        except Exception as e:
            failed += 1
            log(f"ERROR for rank {rank}: {e}")

    cleanup_aux()
    log("=" * 56)
    log("RESUME GENERATION FINISHED")
    log(f"Total jobs: {total}")
    log(f"PDFs generated: {generated}")
    log(f"Already up-to-date: {skipped}")
    log(f"Failed: {failed}")
    log(f"PDF directory: {PDFS_DIR}")
    log("=" * 56)

if __name__ == "__main__":
    main()
